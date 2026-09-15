import openpyxl
import os
import re

try:
    from rapidfuzz import process, fuzz
    HAS_RAPIDFUZZ = True
except ImportError:
    import difflib
    HAS_RAPIDFUZZ = False

class CompanyMatcher:
    def __init__(self, excel_path=None, gspread_worksheet=None):
        self.companies_data = []  # List of dicts: {'Company': ..., 'Role': ...}
        self.unique_company_names = []
        if gspread_worksheet:
            self.load_from_gspread(gspread_worksheet)
        elif excel_path and os.path.exists(excel_path):
            self.load_from_excel(excel_path)

    def load_from_excel(self, excel_path):
        """Load company list and roles from Companies sheet in Excel."""
        self.companies_data = []
        self.unique_company_names = []
        try:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            if 'Companies' in wb.sheetnames:
                ws = wb['Companies']
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row and row[0]:
                        c_name = str(row[0]).strip()
                        c_role = str(row[3]).strip() if len(row) > 3 and row[3] else ""
                        self.companies_data.append({'Company': c_name, 'Role': c_role})
                        if c_name not in self.unique_company_names:
                            self.unique_company_names.append(c_name)
            wb.close()
        except Exception as e:
            print("Error loading companies from Excel for matcher:", e)

    def load_from_gspread(self, gspread_worksheet):
        """Load company list and roles directly from Google Sheets worksheet."""
        self.companies_data = []
        self.unique_company_names = []
        try:
            records = gspread_worksheet.get_all_values()
            if len(records) > 1:
                # header is row 0: Company is col 0, Role is col 3
                for row in records[1:]:
                    if row and len(row) > 0 and row[0].strip():
                        c_name = str(row[0]).strip()
                        c_role = str(row[3]).strip() if len(row) > 3 and row[3] else ""
                        self.companies_data.append({'Company': c_name, 'Role': c_role})
                        if c_name not in self.unique_company_names:
                            self.unique_company_names.append(c_name)
        except Exception as e:
            print("Error loading companies from Google Sheets for matcher:", e)

    def get_exact_company_name(self, input_company_name):
        """
        Finds and returns the EXACT string name as stored in the Companies sheet.
        If 'input_company_name' matches an existing company (exact, case-insensitive, or fuzzy),
        it returns the exact registered company name from the sheet.
        """
        if not input_company_name or not self.unique_company_names:
            return input_company_name.strip() if input_company_name else ""

        clean_input = input_company_name.strip()

        # 1. Exact case-insensitive match
        for existing in self.unique_company_names:
            if existing.lower() == clean_input.lower():
                return existing  # Return exact case from sheet

        # 2. Substring / Prefix match (e.g. "NAVI Tech" -> "NAVI", "Unify Apps" -> "UnifyApps")
        for existing in self.unique_company_names:
            clean_ex = existing.lower().replace(" ", "")
            clean_in = clean_input.lower().replace(" ", "")
            if clean_ex == clean_in or clean_ex in clean_in or clean_in in clean_ex:
                return existing

        # 3. Fuzzy match
        if HAS_RAPIDFUZZ:
            match = process.extractOne(clean_input, self.unique_company_names, scorer=fuzz.WRatio)
            if match and match[1] >= 75:
                return match[0]
        else:
            matches = difflib.get_close_matches(clean_input, self.unique_company_names, n=1, cutoff=0.65)
            if matches:
                return matches[0]

        return clean_input  # Fallback if brand new company

    def find_roles_for_company(self, company_name):
        """Find all roles registered for a company in Companies sheet."""
        exact_name = self.get_exact_company_name(company_name)
        roles = []
        for item in self.companies_data:
            if item['Company'].lower() == exact_name.lower() and item['Role']:
                if item['Role'] not in roles:
                    roles.append(item['Role'])
        return roles
