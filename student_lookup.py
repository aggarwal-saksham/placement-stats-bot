import openpyxl
import os

try:
    from rapidfuzz import process, fuzz
    HAS_RAPIDFUZZ = True
except ImportError:
    import difflib
    HAS_RAPIDFUZZ = False

class StudentLookup:
    def __init__(self, excel_path=None, gspread_worksheet=None):
        self.students = []  # List of dicts: {'Roll No': ..., 'Name': ..., 'CGPA': ..., 'Branch': ...}
        if gspread_worksheet:
            self.load_from_gspread(gspread_worksheet)
        elif excel_path and os.path.exists(excel_path):
            self.load_from_excel(excel_path)

    def load_from_excel(self, excel_path):
        """Load CGPA_Master sheet into memory for fast lookup."""
        self.students = []
        try:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            if 'CGPA_Master' in wb.sheetnames:
                ws = wb['CGPA_Master']
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row and row[0]:
                        roll = str(row[0]).strip().upper()
                        name = str(row[1]).strip().upper() if len(row) > 1 and row[1] else ""
                        cgpa = float(row[2]) if len(row) > 2 and row[2] is not None else None
                        branch = str(row[3]).strip().upper() if len(row) > 3 and row[3] else ""
                        self.students.append({
                            'Roll No': roll,
                            'Name': name,
                            'CGPA': cgpa,
                            'Branch': branch
                        })
            wb.close()
        except Exception as e:
            print("Error loading CGPA_Master from Excel:", e)

    def load_from_gspread(self, gspread_worksheet):
        """Load CGPA_Master sheet directly from Google Sheets."""
        self.students = []
        try:
            records = gspread_worksheet.get_all_values()
            if len(records) > 1:
                # row 0 is header: Roll No, Name, CGPA, Branch
                for row in records[1:]:
                    if row and len(row) > 0 and row[0].strip():
                        roll = str(row[0]).strip().upper()
                        name = str(row[1]).strip().upper() if len(row) > 1 and row[1] else ""
                        branch = str(row[3]).strip().upper() if len(row) > 3 and row[3] else ""
                        self.students.append({
                            'Roll No': roll,
                            'Name': name,
                            'CGPA': None,
                            'Branch': branch
                        })
        except Exception as e:
            print("Error loading CGPA_Master from Google Sheets:", e)

    def find_by_roll(self, roll_no):
        """Find student by Roll No (exact or normalized)."""
        roll_clean = roll_no.strip().upper()
        for s in self.students:
            if s['Roll No'] == roll_clean:
                return s
        return None

    def find_by_name(self, name):
        """Find student by Name using exact or fuzzy match."""
        name_clean = name.strip().upper()
        # 1. Exact match
        for s in self.students:
            if s['Name'] == name_clean:
                return s, 100.0

        # 2. Fuzzy match
        all_names = [s['Name'] for s in self.students]
        if not all_names:
            return None, 0.0

        if HAS_RAPIDFUZZ:
            match = process.extractOne(name_clean, all_names, scorer=fuzz.WRatio)
            if match and match[1] >= 85:
                for s in self.students:
                    if s['Name'] == match[0]:
                        return s, match[1]
        else:
            matches = difflib.get_close_matches(name_clean, all_names, n=1, cutoff=0.75)
            if matches:
                for s in self.students:
                    if s['Name'] == matches[0]:
                        return s, 90.0

        return None, 0.0

    def enrich_student_data(self, input_roll="", input_name=""):
        """
        Given partial input (Roll No or Name), auto-fills missing fields from CGPA_Master.
        """
        result = {
            'Roll No': input_roll.strip().upper() if input_roll else "",
            'Name': input_name.strip().upper() if input_name else ""
        }

        # Case 1: Roll No provided
        if result['Roll No']:
            match = self.find_by_roll(result['Roll No'])
            if match:
                if not result['Name']:
                    result['Name'] = match['Name']
                return result

        # Case 2: Name provided
        if result['Name']:
            match, score = self.find_by_name(result['Name'])
            if match and score >= 75:
                if not result['Roll No']:
                    result['Roll No'] = match['Roll No']
                result['Name'] = match['Name']  # Normalize name

        return result
